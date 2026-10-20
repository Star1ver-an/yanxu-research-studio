import base64
import io
import json
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
import queue
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.request import Request,urlopen
from urllib.error import HTTPError
import zipfile

from app.storage import Store,extract_file,protect
from app.service import Service,make_docx
from app.providers import CodexClient,ProviderError,Cancelled,api_complete,validate_endpoint
from app.server import LocalServer

class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.store=Store(self.temp.name)
        self.service=Service(self.store)
        self.project=self.store.create({'title':'验证项目','flow':'full','experiments':'none'})
        self.pid=self.project['id']
    def tearDown(self):
        self.temp.cleanup()
    def output(self,sid,text='可检查的报告'):
        p=self.store.get(self.pid)
        p['outputs'][sid]={'text':text,'approved':False,'stale':False}
        self.store.put(p)
    def test_create_validation_and_persistence(self):
        with self.assertRaises(ValueError):
            self.store.create({'title':'','flow':'full'})
        self.assertEqual(Store(self.temp.name).get(self.pid)['title'],'验证项目')
    def test_cannot_confirm_empty_or_out_of_order(self):
        with self.assertRaises(ValueError):
            self.service.confirm(self.pid,{'stage':'research'})
        self.output('outline')
        with self.assertRaises(ValueError):
            self.service.confirm(self.pid,{'stage':'outline'})
    def test_confirm_records_hash_and_advances_once(self):
        self.output('research')
        p=self.service.confirm(self.pid,{'stage':'research'})
        self.assertEqual(p['cursor'],1)
        self.assertEqual(len(p['confirmations'][0]['sha256']),64)
        with self.assertRaises(ValueError):
            self.service.confirm(self.pid,{'stage':'research'})
    def test_gate_requires_explicit_acknowledgment(self):
        p=self.store.get(self.pid);p['cursor']=3;self.store.put(p)
        self.output('integrity','核验范围有限，关键论断 UNKNOWN')
        with self.assertRaises(ValueError):
            self.service.confirm(self.pid,{'stage':'integrity'})
        self.assertEqual(self.service.confirm(self.pid,{'stage':'integrity','acknowledge':True})['cursor'],4)
    def test_undecided_experiment_declaration_stops_advancement(self):
        p=self.store.get(self.pid);p['experiments']='undecided';self.store.put(p)
        self.output('research')
        with self.assertRaises(ValueError):
            self.service.confirm(self.pid,{'stage':'research'})
    def test_edit_invalidates_downstream_without_erasing_work(self):
        p=self.store.get(self.pid);p['cursor']=4
        p['outputs']={k:{'text':k,'approved':True,'stale':False} for k in ['research','outline','draft','integrity']}
        self.store.put(p)
        p=self.service.edit(self.pid,{'stage':'outline','text':'更新过的提纲'})
        self.assertEqual(p['cursor'],1)
        self.assertTrue(p['outputs']['draft']['stale'])
        self.assertEqual(p['outputs']['draft']['text'],'draft')
        self.assertTrue(p['outputs']['research']['approved'])
    def test_materials_are_parsed_and_private_text_not_sent_to_ui(self):
        self.output('research')
        p=self.service.upload(self.pid,{'name':'../../../记录.md','data':base64.b64encode('真实材料'.encode()).decode()})
        self.assertEqual(p['assets'][0]['name'],'记录.md')
        self.assertNotIn('text',p['assets'][0])
        self.assertTrue(p['outputs']['research']['stale'])
        full=self.store.get(self.pid)
        self.assertEqual(full['assets'][0]['text'],'真实材料')
        self.assertTrue((self.store.root/full['assets'][0]['file']).exists())
    def test_context_refuses_oversize_instead_of_silent_truncation(self):
        p=self.store.get(self.pid)
        p['assets']=[{'enabled':True,'name':'big.txt','text':'a'*240001}]
        with self.assertRaisesRegex(ValueError,'24 万'):
            self.service.context(p,'research','')
    def test_input_changes_invalidate_confirmed_outputs(self):
        self.output('research')
        self.service.confirm(self.pid,{'stage':'research'})
        p=self.service.update_info(self.pid,{'topic':'不同的研究目标'})
        self.assertEqual(p['cursor'],0)
        self.assertTrue(p['outputs']['research']['stale'])
    def test_restart_marks_incomplete_job_interrupted(self):
        p=self.store.get(self.pid);p['runs']=[{'id':'r1','status':'running'}];self.store.put(p)
        recovered=Store(self.temp.name).get(self.pid)
        self.assertEqual(recovered['runs'][0]['status'],'interrupted')
    def test_export_docx_and_backup_are_readable(self):
        self.output('research','# 报告\n\n| 列 A | 列 B |\n|---|---|\n| 证据 | 未知 |\n\n- 待核实')
        name,raw,mime=self.service.export(self.pid,'docx')
        from docx import Document
        d=Document(io.BytesIO(raw))
        self.assertTrue(any('报告' in p.text for p in d.paragraphs))
        self.assertEqual(d.tables[0].cell(1,0).text,'证据')
        name,raw,mime=self.service.export(self.pid,'zip')
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            data=json.loads(z.read('项目备份.json'))
            self.assertEqual(data['project']['title'],'验证项目')
            self.assertNotIn('api_key',z.read('项目备份.json').decode())
    def test_dpapi_roundtrip_and_no_key_echo(self):
        self.assertEqual(protect(protect('test-secret'),True),'test-secret')
        s=self.store.save_settings({'api_key':'test-secret','mode':'api'})
        self.assertTrue(s['has_key'])
        self.assertNotIn('test-secret',json.dumps(s))
        self.assertEqual(self.store.get_settings(secret=True)['api_key'],'test-secret')
        self.assertNotIn(b'test-secret',self.store.db.read_bytes())
    def test_async_execution_saves_unconfirmed_result(self):
        class FakeCodex:
            def __init__(self,*a):pass
            def __enter__(self):return self
            def __exit__(self,*a):pass
            def complete(self,*a):return '# 运行成功\n\n材料不足，等待作者确认。'
        with patch('app.service.CodexClient',FakeCodex):
            p=self.service.start(self.pid,{'stage':'research','note':'仅用材料'})
            for _ in range(100):
                p=self.store.get(self.pid)
                if p['runs'][-1]['status']!='running':break
                time.sleep(.01)
        self.assertEqual(p['runs'][-1]['status'],'completed')
        self.assertFalse(p['outputs']['research']['approved'])
        self.assertEqual(p['cursor'],0)
    def test_failed_execution_preserves_old_report(self):
        self.output('research','已有报告')
        class FailedCodex:
            def __init__(self,*a):raise ProviderError('连接不可用')
        with patch('app.service.CodexClient',FailedCodex):
            self.service.start(self.pid,{'stage':'research'})
            for _ in range(100):
                p=self.store.get(self.pid)
                if p['runs'][-1]['status']!='running':break
                time.sleep(.01)
        self.assertEqual(p['runs'][-1]['status'],'failed')
        self.assertEqual(p['outputs']['research']['text'],'已有报告')
    def test_api_payloads_match_both_protocols(self):
        received=[]
        class API(BaseHTTPRequestHandler):
            def log_message(self,*a):pass
            def do_POST(self):
                data=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                received.append((self.path,dict(self.headers),data))
                response={'content':[{'type':'text','text':'测试成功'}],'stop_reason':'end_turn'} if self.path.endswith('messages') else {'choices':[{'message':{'content':'测试成功'},'finish_reason':'stop'}]}
                raw=json.dumps(response).encode();self.send_response(200);self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        server=ThreadingHTTPServer(('127.0.0.1',0),API)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            for provider in ['compatible','anthropic']:
                s={'provider':provider,'base_url':f'http://127.0.0.1:{server.server_port}/v1','api_model':'test-model'}
                self.assertEqual(api_complete(s,'test-key','system','prompt',lambda *a:None,threading.Event()),'测试成功')
            self.assertEqual(received[0][0],'/v1/chat/completions')
            self.assertEqual(received[1][0],'/v1/messages')
            self.assertEqual(received[1][2]['system'],'system')
            self.assertEqual(received[0][2]['messages'][0]['role'],'system')
        finally:
            server.shutdown();server.server_close()
    def test_remote_plain_http_and_malformed_urls_rejected(self):
        for url in ['file:///a','http://remote.example/v1','https://name:secret@example.com']:
            with self.assertRaises(ProviderError):validate_endpoint(url)
    def test_local_server_checks_token_and_host(self):
        server=LocalServer(self.service)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        base=f'http://127.0.0.1:{server.server_port}'
        try:
            with self.assertRaises(HTTPError):urlopen(base+'/api/bootstrap')
            with urlopen(Request(base+'/api/bootstrap',headers={'X-Research-Token':server.token})) as r:
                self.assertEqual(json.load(r)['projects'][0]['id'],self.pid)
            with self.assertRaises(HTTPError):urlopen(Request(base,headers={'Host':'evil.example'}))
        finally:
            server.shutdown();server.server_close()
    def test_supported_document_parsers(self):
        from docx import Document
        d=Document();d.add_paragraph('真实稿件文本');raw=io.BytesIO();d.save(raw)
        self.assertIn('真实稿件文本',extract_file('paper.docx',raw.getvalue()))
        self.assertEqual(extract_file('中文.txt','文献记录'.encode('gb18030')),'文献记录')
        with self.assertRaises(ValueError):extract_file('program.exe',b'bad')

    def test_desktop_bridge_exposes_only_intended_methods(self):
        from main import DesktopAPI
        bridge=DesktopAPI(self.service)
        public=[name for name in dir(bridge) if not name.startswith('_')]
        self.assertEqual(set(public),{'save_export','save_assistant_export','save_progress_export','open_data_folder','open_link','report_ready'})
        self.assertTrue(all(callable(getattr(bridge,name)) for name in public))

    def test_native_policy_supports_official_bridge_wrapper_factory(self):
        server=LocalServer(self.service)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        base=f'http://127.0.0.1:{server.server_port}'
        try:
            with urlopen(base) as r:
                self.assertNotIn('unsafe-eval',r.headers['Content-Security-Policy'])
            server.native=object()
            with urlopen(base) as r:
                self.assertIn("script-src 'self' 'unsafe-eval'",r.headers['Content-Security-Policy'])
        finally:
            server.shutdown();server.server_close()

    def test_desktop_ready_state_cannot_be_overwritten_by_load_event(self):
        from main import DesktopAPI
        bridge=DesktopAPI(self.service)
        bridge.report_ready()
        bridge._record_loaded()
        state=json.loads((self.store.root/'runtime.json').read_text())
        self.assertTrue(bridge._ready)
        self.assertEqual(state['state'],'ui_ready')

if __name__=='__main__':
    unittest.main()
