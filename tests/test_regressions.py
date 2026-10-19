"""Regression cases for issues found through the user-facing workflow."""
import base64
import io
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import queue
import tempfile
import threading
import unittest
from unittest.mock import patch

from app.literature import filter_entries,identifier_terms,metadata_matches
from app.providers import ReplyCollector,CodexClient,ProviderError,Cancelled,api_complete
from app.service import Service,make_docx
from app.storage import Store


class MessageTests(unittest.TestCase):
    def test_commentary_never_becomes_preview_or_report(self):
        c=ReplyCollector()
        c.item(dict(type='agentMessage',id='c',phase='commentary',text='我要先检查材料'),True)
        c.item(dict(type='agentMessage',id='f',phase='final_answer',text=''))
        c.delta('f','# 报告\n正在生成蛋白的过程值得研究。')
        self.assertNotIn('我要先',c.preview())
        c.item(dict(type='agentMessage',id='f',phase='final_answer',text='# 最终报告\n正在生成蛋白。'),True)
        self.assertEqual(c.result(),'# 最终报告\n正在生成蛋白。')
        self.assertEqual(c.metadata()['message_phases'],{'commentary':1,'final_answer':1})

    def test_phase_can_arrive_only_on_completion(self):
        c=ReplyCollector();c.delta('c','进度');c.delta('f','草稿')
        c.item(dict(type='agentMessage',id='c',phase='commentary',text='进度'),True)
        c.item(dict(type='agentMessage',id='f',phase='final_answer',text='正文'),True)
        self.assertEqual(c.result(),'正文')

    def test_legacy_keeps_last_message_instead_of_all_preambles(self):
        c=ReplyCollector()
        for key,text in [('1','准备'),('2','正文')]:
            c.item(dict(type='agentMessage',id=key,text=text),True)
        self.assertEqual(c.result(),'正文')
        self.assertTrue(c.metadata()['legacy_fallback'])

    def test_commentary_or_partial_only_is_not_a_complete_report(self):
        for phase in ['commentary','partial_answer']:
            c=ReplyCollector();c.item(dict(type='agentMessage',id='1',phase=phase,text='未结束'),True)
            with self.assertRaises(ProviderError):c.result()

    def test_full_event_loop_filters_other_turn_and_uses_completed_item(self):
        c=CodexClient.__new__(CodexClient)
        c.settings={};c._cwd='.';c.saved_events=[]
        c.account=lambda:dict(type='chatgpt')
        c.request=lambda method,*a,**kw: {'thread':{'id':'th'}} if method=='thread/start' else {'turn':{'id':'tu'}}
        events=iter([
          {'method':'item/started','params':{'item':dict(type='agentMessage',id='1',phase='commentary',text='')}},
          {'method':'item/agentMessage/delta','params':{'itemId':'1','delta':'先检查'}},
          {'method':'item/completed','params':{'item':dict(type='agentMessage',id='1',phase='commentary',text='先检查')}},
          {'method':'item/completed','params':{'turnId':'other','item':dict(type='agentMessage',id='x',phase='final_answer',text='别的任务')}},
          {'method':'item/started','params':{'item':dict(type='agentMessage',id='2',phase='final_answer',text='')}},
          {'method':'item/agentMessage/delta','params':{'itemId':'2','delta':'草稿'}},
          {'method':'item/completed','params':{'item':dict(type='agentMessage',id='2',phase='final_answer',text='完成正文')}},
          {'method':'turn/completed','params':{'turn':{'id':'tu','status':'completed'}}}])
        c.next=lambda *a:next(events)
        previews=[]
        self.assertEqual(c.complete('system','prompt','',lambda *a:previews.append(a)),'完成正文')
        self.assertFalse(any('先检查' in x[-1] for x in previews))


class LiteratureTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.store=Store(self.temp.name);self.service=Service(self.store)
        self.p=self.store.create({'title':'TMEM43 课题','flow':'full','experiments':'none'})
    def tearDown(self):self.temp.cleanup()

    def test_identifier_boundaries_and_chinese(self):
        self.assertEqual(identifier_terms('TMEM43 renal IRI-AKI TMEM43'),['TMEM43'])
        self.assertFalse(metadata_matches({'title':'TMEM430 study'},['TMEM43']))
        self.assertTrue(metadata_matches({'title':'TMEM43在肾损伤中的作用'},['TMEM43']))
        self.assertTrue(metadata_matches({'title':'研究缺血再灌注损伤'},['缺血再灌注']))

    def test_record_type_and_supplements_are_transparent(self):
        entries=[dict(title='TMEM43',type=t,is_supplement=s,doi=str(i))
                 for i,(t,s) in enumerate([('journal-article',False),('proceedings-article',False),('journal-article',True)])]
        kept,rejected=filter_entries(entries,['TMEM43'],True)
        self.assertEqual(len(kept),1);self.assertEqual(len(rejected),2)
        self.assertEqual(len(filter_entries(entries,[],False)[0]),3)
        supplement=dict(title='TMEM43 meeting abstract',type='journal-article',issue='Supplement_1')
        self.assertEqual(filter_entries([supplement],[],True)[0],[])

    def test_search_is_unselected_deduplicated_and_does_not_invalidate(self):
        p=self.store.get(self.p['id']);p['cursor']=1;p['outputs']['research']=dict(text='既有成果',approved=True,stale=False);self.store.put(p)
        records=[{'DOI':'10.1/x','title':['TMEM43 renal study'],'type':'journal-article'},
                 {'DOI':'10.1/X','title':['TMEM43 duplicate'],'type':'journal-article'},
                 {'DOI':'10.1/y','title':['Unrelated renal study'],'type':'journal-article'}]
        response=io.BytesIO(json.dumps({'message':{'items':records}}).encode())
        with patch('app.service.urlopen',return_value=response):r=self.service.search(p['id'],'TMEM43 kidney')
        self.assertEqual(r['added'],1);self.assertEqual(len(r['rejected']),1)
        p=self.store.get(p['id']);self.assertEqual(p['cursor'],1);self.assertTrue(p['outputs']['research']['approved'])
        self.assertNotIn('10.1/x',self.service.context(p,'research',''))
        self.service.toggle_literature(p['id'],dict(id='10.1/x',enabled=True))
        p=self.store.get(p['id']);self.assertIn('10.1/x',self.service.context(p,'research',''))
        self.assertTrue(p['outputs']['research']['stale']);self.assertEqual(p['cursor'],0)
        self.service.toggle_literature(p['id'],dict(all=True,enabled=False))
        self.assertFalse(self.store.get(p['id'])['literature'][0]['enabled'])

    def test_old_records_remain_enabled_until_user_changes_selection(self):
        p=self.store.get(self.p['id']);p['literature']=[dict(id='old',title='保留旧记录')]
        self.assertIn('保留旧记录',self.service.context(p,'research',''))

    def test_upload_server_rejects_oversize_without_writing(self):
        raw=base64.b64encode(b'x'*(20*1024*1024+1)).decode()
        with self.assertRaisesRegex(ValueError,'20 MB'):
            self.service.upload(self.p['id'],dict(name='大文件.txt',data=raw))
        self.assertEqual(self.store.get(self.p['id'])['assets'],[])

    def test_confirm_records_delegate_without_claiming_author(self):
        p=self.store.get(self.p['id']);p['outputs']['research']=dict(text='报告',approved=False,stale=False);self.store.put(p)
        p=self.service.confirm(p['id'],dict(stage='research',actor='delegate'))
        self.assertEqual(p['confirmations'][0]['actor'],'delegate')
        self.assertNotIn('作者确认',p['confirmations'][0]['meaning'])


class APIIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.received=[]
        class Mock(BaseHTTPRequestHandler):
            def log_message(self,*a):pass
            def do_POST(self):
                data=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                cls.received.append((self.path,dict(self.headers),data))
                model=data['model'];status=int(model[4:]) if model.startswith('http') else 200
                if model=='invalid-json':raw=b'not json'
                elif model=='array':raw=b'[]'
                elif status!=200:raw=b'{"error":"private-server-body test-key"}'
                else:
                    text='' if model=='empty' else '本机模拟正文'
                    if self.path.endswith('/messages'):
                        body={'content':[{'type':'thinking','thinking':'内部推理'},{'type':'text','text':text}],
                              'stop_reason':'max_tokens' if model=='truncated' else 'end_turn'}
                    else:
                        content=[{'type':'reasoning','text':'内部推理'},{'type':'text','text':text}] if model=='blocks' else text
                        body={'choices':[{'message':{'content':content},'finish_reason':'length' if model=='truncated' else 'stop'}]}
                    raw=json.dumps(body,ensure_ascii=False).encode()
                self.send_response(status);self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),Mock)
        threading.Thread(target=cls.server.serve_forever,daemon=True).start()
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close()
    def settings(self,provider,model='success'):
        return dict(provider=provider,base_url=f'http://127.0.0.1:{self.server.server_port}/v1',api_model=model,max_tokens=1234)
    def complete(self,provider='compatible',model='success',cancel=None,progress=None):
        return api_complete(self.settings(provider,model),'test-key','system','prompt',progress or (lambda *a:None),cancel or threading.Event())
    def test_success_and_authentication_for_both_protocols(self):
        for provider,path in [('compatible','/v1/chat/completions'),('anthropic','/v1/messages')]:
            self.assertEqual(self.complete(provider),'本机模拟正文')
            route,headers,data=self.received[-1]
            self.assertEqual(route,path);self.assertEqual(data['max_tokens'],1234)
            headers={k.lower():v for k,v in headers.items()}
            self.assertEqual(headers['authorization' if provider=='compatible' else 'x-api-key'],'Bearer test-key' if provider=='compatible' else 'test-key')
    def test_status_errors_are_specific_and_do_not_echo_server_body(self):
        for provider in ['compatible','anthropic']:
            for code in [401,403,404,429,500]:
                with self.assertRaises(ProviderError) as ctx:self.complete(provider,'http'+str(code))
                self.assertIn(str(code),str(ctx.exception));self.assertNotIn('test-key',str(ctx.exception));self.assertNotIn('private-server',str(ctx.exception))
    def test_invalid_json_empty_and_wrong_shapes(self):
        for provider in ['compatible','anthropic']:
            for model in ['invalid-json','empty','array']:
                with self.assertRaises(ProviderError):self.complete(provider,model)
    def test_length_limit_is_visible_and_reasoning_blocks_are_excluded(self):
        for provider in ['compatible','anthropic']:
            self.assertIn('输出达到长度上限',self.complete(provider,'truncated'))
        self.assertEqual(self.complete(model='blocks'),'本机模拟正文')
    def test_cancellation_before_and_after_response(self):
        event=threading.Event();event.set();count=len(self.received)
        with self.assertRaises(Cancelled):self.complete(cancel=event)
        self.assertEqual(len(self.received),count)
        event.clear()
        with patch('app.providers.json.loads',side_effect=lambda b:(event.set(),json.JSONDecoder().decode(b.decode()))[1]):
            with self.assertRaises(Cancelled):self.complete(cancel=event)


class WordLayoutTests(unittest.TestCase):
    def test_a4_title_pagination_headers_and_inline_formatting(self):
        from docx import Document
        from docx.oxml.ns import qn
        doc=Document(io.BytesIO(make_docx('# 研究方案\n\n## 第一节\n正文 *斜体* **加粗** `代码`\n\n| 组别 | 方案 |\n|---|---|\n| 对照 | 待执行 |')))
        self.assertAlmostEqual(doc.sections[0].page_width.cm,21,places=1)
        self.assertEqual(doc.paragraphs[0].style.name,'Title')
        self.assertIsNone(doc.styles['Title'].element.pPr.find(qn('w:pBdr')))
        self.assertTrue(doc.styles['Heading 2'].paragraph_format.keep_with_next)
        runs=[r for p in doc.paragraphs for r in p.runs]
        self.assertTrue(any(r.italic and r.text=='斜体' for r in runs))
        self.assertTrue(any(r.bold and r.text=='加粗' for r in runs))
        self.assertIsNotNone(doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')))
        self.assertIsNotNone(doc.tables[0].rows[1]._tr.trPr.find(qn('w:cantSplit')))


if __name__=='__main__':unittest.main()
