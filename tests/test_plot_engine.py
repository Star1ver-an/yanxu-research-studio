"""Actual bounded-reader and renderer checks, using synthetic nonresearch data."""
import copy
import hashlib
import io
import unittest

from app.plot_data import inspect_source, load_table, MAX_BYTES
from app.plot_engine import catalog, render_figure, validate_spec


def source(text, name='evaluation.csv', ident='table'):
    return {'id': ident, 'name': name, 'raw': text.encode('utf-8')}


def recipe(kind, mapping=None, **extra):
    return {'version': 1, 'kind': kind, 'title': '合成 AI 测试：训练与评估', 'source_ids': ['table'],
            'mapping': mapping or {}, 'style': {'width': 6, 'height': 4, 'dpi': 90}, **extra}


class PlotDataTests(unittest.TestCase):
    def test_utf8_missing_and_preview_are_bounded(self):
        raw = ('epoch,loss,note\n' + '\n'.join(f'{i},{i / 10},记录{i}' for i in range(12))).encode()
        details = inspect_source(raw, 'test.csv')
        self.assertEqual(details['row_count'], 12)
        self.assertEqual(len(details['preview']), 8)
        self.assertEqual(details['columns'][1]['type'], 'numeric')
        self.assertEqual(details['sha256'], hashlib.sha256(raw).hexdigest())
        table = load_table(b'x,y\n1,NaN\n2,Infinity\n3,NA\n4,0', 'x.csv')
        self.assertEqual(table['columns'][1]['missing'], 3)
        self.assertEqual(table['rows'][-1]['y'], 0)

    def test_identifiers_and_large_integer_strings_are_preserved(self):
        table = load_table(b'id,value\n001,2\n12345678901234567890,3', 'id.csv')
        self.assertEqual(table['rows'][0]['id'], '001')
        self.assertEqual(table['rows'][1]['id'], '12345678901234567890')

    def test_tsv_and_legacy_encoding(self):
        info = inspect_source('模型\t准确率\n轻量模型\t0.9'.encode('gb18030'), 'x.tsv')
        self.assertEqual(info['preview'][0]['模型'], '轻量模型')
        self.assertTrue(info['warnings'])

    def test_malformed_and_oversize_inputs_fail_clearly(self):
        for raw, name, message in [(b'a,a\n1,2', 'x.csv', '重复'), (b'a,b\n1,2,3', 'x.csv', '不一致'),
                                   (b'a,b\n"unterminated', 'x.csv', '格式'), (b'abc', 'x.exe', '支持'),
                                   (b'x' * (MAX_BYTES + 1), 'x.csv', '20 MB')]:
            with self.subTest(name=name, message=message), self.assertRaisesRegex(ValueError, message):
                inspect_source(raw, name)

    def test_xlsx_values_and_first_sheet_notice(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        workbook.active.append(['模型', '准确率'])
        workbook.active.append(['tiny', .9])
        workbook.create_sheet('other')
        buf = io.BytesIO();workbook.save(buf)
        info = inspect_source(buf.getvalue(), 'scores.xlsx')
        self.assertEqual(info['preview'][0]['准确率'], .9)
        self.assertIn('第一张', info['warnings'][0])

    def test_excel_formula_is_not_executed_or_trusted(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        workbook.active.append(['epoch', 'score'])
        workbook.active.append([1, '=WEBSERVICE("https://invalid.example")'])
        buf = io.BytesIO();workbook.save(buf)
        with self.assertRaisesRegex(ValueError, '公式'):
            inspect_source(buf.getvalue(), 'formula.xlsx')

    def test_png_source_dimensions_and_bytes_remain_unchanged(self):
        from PIL import Image
        output = io.BytesIO();Image.new('RGB', (80, 50), '#385E85').save(output, format='PNG')
        raw = output.getvalue();before = hashlib.sha256(raw).hexdigest()
        info = inspect_source(raw, 'source.png')
        self.assertEqual((info['width'], info['height']), (80, 50))
        self.assertEqual(info['sha256'], before)


class PlotEngineTests(unittest.TestCase):
    def render(self, spec, inputs, fmt='png'):
        before = copy.deepcopy(spec)
        result = render_figure(spec, inputs, fmt)
        self.assertEqual(spec, before, 'Rendering must not mutate recipes')
        self.assertGreater(len(result['data']), 1000)
        self.assertFalse(result['metadata']['model_called'])
        self.assertEqual(result['metadata']['sources'][0]['sha256'], hashlib.sha256(inputs[0]['raw']).hexdigest())
        return result

    def test_catalog_does_not_claim_specialist_analysis(self):
        items = catalog()
        self.assertEqual(len(items), 18)
        self.assertEqual(len({i['id'] for i in items}), 18)
        self.assertTrue(all(i['support'] == 'results' for i in items if i['id'] in {'km', 'forest', 'roc', 'volcano'}))
        items[0]['label'] = 'changed'
        self.assertNotEqual(catalog()[0]['label'], 'changed')

    def test_general_chart_kinds_with_real_rendering(self):
        data = source('epoch,loss,group,model\n1,.9,A,Tiny\n2,.6,A,Base\n3,.4,A,Large\n1,.8,B,Tiny\n2,.5,B,Base\n3,.3,B,Large')
        for kind, mapping in [('line', {'x': 'epoch', 'y': 'loss', 'group': 'group'}),
                              ('scatter', {'x': 'epoch', 'y': 'loss', 'group': 'group'}),
                              ('bar', {'x': 'model', 'y': 'loss', 'group': 'group'}),
                              ('box', {'x': 'group', 'y': 'loss'}), ('violin', {'x': 'group', 'y': 'loss'}),
                              ('histogram', {'x': 'loss', 'group': 'group'})]:
            with self.subTest(kind=kind):
                result = self.render(recipe(kind, mapping), [data])
                self.assertEqual(result['data'][:8], b'\x89PNG\r\n\x1a\n')
                self.assertEqual(result['metadata']['rows_used'], 6)

    def test_wide_multiline_and_provenance(self):
        spec = recipe('line', {'x': 'epoch', 'y': ['train', 'valid']})
        result = self.render(spec, [source('epoch,train,valid\n1,.8,.9\n2,.4,.5')])
        self.assertEqual(result['metadata']['rows_used'], 2)
        self.assertTrue(result['metadata']['spec_sha256'])

    def test_specialist_results_render_without_inventing_statistics(self):
        cases = [
            ('volcano', 'fc,p,name\n-2,.001,A\n.1,.8,B\n3,.01,C', {'x': 'fc', 'y': 'p'}),
            ('ma', 'mean,fc\n2,-1\n4,.2\n6,2', {'x': 'mean', 'y': 'fc'}),
            ('forest', 'model,effect,lo,hi\nTiny,.8,.6,1.1\nBase,1.2,1.1,1.4', {'y': 'effect', 'lower': 'lo', 'upper': 'hi', 'label': 'model'}),
            ('roc', 'fpr,tpr\n0,0\n.1,.6\n.5,.9\n1,1', {'x': 'fpr', 'y': 'tpr'}),
            ('pr', 'recall,precision\n0,1\n.5,.8\n1,.3', {'x': 'recall', 'y': 'precision'}),
            ('km', 'time,survival\n0,1\n1,.9\n2,.7', {'x': 'time', 'y': 'survival'}),
            ('enrichment', 'term,ratio,count,p\nPathA,.2,8,.01\nPathB,.1,4,.04', {'x': 'ratio', 'y': 'term', 'size': 'count', 'color': 'p'}),
        ]
        for kind, csv, mapping in cases:
            with self.subTest(kind=kind):
                result = self.render(recipe(kind, mapping), [source(csv)])
                self.assertFalse(result['metadata']['analysis_performed'])

    def test_heatmap_and_confusion_matrix(self):
        data = source('actual,tiny,base\ntiny,8,2\nbase,1,9')
        for kind in ['heatmap', 'confusion']:
            result = self.render(recipe(kind, {'x': 'actual', 'y': ['tiny', 'base']}), [data])
            self.assertEqual(result['metadata']['matrix_shape'], [2, 2])
            self.assertEqual(result['metadata']['transforms'], [])

    def test_long_heatmap_never_aggregates_duplicate_values(self):
        spec = recipe('heatmap', {'x': 'col', 'y': 'row', 'color': 'value'})
        result = self.render(spec, [source('row,col,value\nA,X,1\nA,Y,2\nB,X,3')])
        self.assertTrue(any('未补零' in warning for warning in result['warnings']))
        with self.assertRaisesRegex(ValueError, '多个值'):
            render_figure(spec, [source('row,col,value\nA,X,1\nA,X,2')])

    def test_missing_rows_are_reported_and_not_imputed(self):
        result = self.render(recipe('scatter', {'x': 'x', 'y': 'y'}), [source('x,y\n1,2\n2,NA\n3,4')])
        self.assertEqual(result['metadata']['rows_used'], 2)
        self.assertTrue(any('排除' in w for w in result['warnings']))

    def test_bar_replicates_require_explicit_mean(self):
        spec = recipe('bar', {'x': 'model', 'y': 'score'})
        data = source('model,score\nTiny,.8\nTiny,.9\nBase,.95')
        with self.assertRaisesRegex(ValueError, 'mean'):
            render_figure(spec, [data])
        spec['options'] = {'aggregate': 'mean'}
        result = self.render(spec, [data])
        self.assertEqual(result['metadata']['aggregation'], 'mean')
        self.assertAlmostEqual(result['metadata']['bar_values'][0]['value'], .85)
        self.assertEqual(result['metadata']['bar_values'][0]['record_count'], 2)
        self.assertTrue(any('未生成误差线' in w for w in result['warnings']))

    def test_invalid_curves_intervals_and_zero_p_fail(self):
        cases = [
            (recipe('km', {'x': 'x', 'y': 'y'}), 'x,y\n0,.7\n1,.9', '不随时间增加'),
            (recipe('roc', {'x': 'x', 'y': 'y'}), 'x,y\n0,0\n1,1.2', '0–1'),
            (recipe('roc', {'x': 'x', 'y': 'y'}), 'x,y\n1,1\n0,0', '递增'),
            (recipe('volcano', {'x': 'x', 'y': 'y'}), 'x,y\n1,0\n2,.1', 'p=0'),
            (recipe('forest', {'y': 'x', 'lower': 'lo', 'upper': 'hi', 'label': 'name'}), 'name,x,lo,hi\nA,2,3,4', '下界'),
        ]
        for spec, csv, message in cases:
            with self.subTest(kind=spec['kind']), self.assertRaisesRegex(ValueError, message):
                render_figure(spec, [source(csv)])

    def test_exports_are_actual_vector_or_raster_formats_and_reproducible(self):
        spec = recipe('scatter', {'x': 'x', 'y': 'y'})
        data = source('x,y\n1,2\n2,3\n3,4')
        signatures = {'png': b'\x89PNG', 'pdf': b'%PDF', 'svg': b'<?xml', 'tiff': b'II'}
        for fmt in ['png', 'svg', 'pdf', 'tiff']:
            with self.subTest(fmt=fmt):
                first = self.render(spec, [data], fmt)
                second = self.render(spec, [data], fmt)
                self.assertTrue(first['data'].startswith(signatures[fmt]))
                self.assertEqual(first['data'], second['data'])
        svg = self.render(spec, [data], 'svg')['data'].decode()
        self.assertNotIn('<script', svg)
        self.assertIn('<path', svg)

    def test_diagram_is_local_editable_and_does_not_require_data(self):
        spec = recipe('diagram', diagram={'nodes': [{'id': 'a', 'label': '原始数据'}, {'id': 'b', 'label': '训练与独立测试'}],
                                          'edges': [{'from': 'a', 'to': 'b', 'label': '数据划分'}]})
        spec['source_ids'] = []
        result = render_figure(spec, [])
        self.assertEqual(result['metadata']['node_count'], 2)
        self.assertEqual(result['metadata']['sources'], [])
        self.assertTrue(result['data'].startswith(b'\x89PNG'))

    def test_panel_preserves_source_and_declares_display_operations(self):
        from PIL import Image
        buf = io.BytesIO();Image.new('RGB', (100, 50), '#BC7052').save(buf, format='PNG')
        image = {'id': 'image', 'name': 'original.png', 'raw': buf.getvalue()}
        before = hashlib.sha256(image['raw']).hexdigest()
        spec = recipe('panel', panels=[{'asset_id': 'image', 'label': 'A · 架构示例'}])
        spec['source_ids'] = ['image']
        result = self.render(spec, [image])
        self.assertEqual(hashlib.sha256(image['raw']).hexdigest(), before)
        self.assertEqual(result['metadata']['image_processing'], ['display_exif_orientation', 'fit_preserve_aspect', 'label_outside_image'])

    def test_schema_resource_limits_and_unsafe_fields_are_rejected(self):
        valid = recipe('scatter', {'x': 'x', 'y': 'y'})
        changes = [({'kind': []}, '图型'), ({'code': 'import os'}, '不支持'),
                   ({'style': {'width': 16, 'height': 16, 'dpi': 600}}, '像素'),
                   ({'style': {'palette': []}}, '配色'), ({'style': {'dpi': float('inf')}}, '非有限'),
                   ({'mapping': {'url': 'https://invalid.example'}}, '不支持'),
                   ({'options': {'javascript': 'alert(1)'}}, '不支持')]
        for patch, message in changes:
            with self.subTest(patch=patch), self.assertRaisesRegex(ValueError, message):
                validate_spec({**valid, **patch})
        with self.assertRaisesRegex(ValueError, '节点'):
            validate_spec(recipe('diagram', diagram={'nodes': [{'id': 'a', 'label': 'x'}], 'edges': [{'from': 'a', 'to': 'missing'}]}))
        with self.assertRaisesRegex(ValueError, '来源'):
            validate_spec(recipe('panel', panels=[{'asset_id': 'missing'}]))

    def test_unknown_or_unselected_source_cannot_be_rendered(self):
        spec = recipe('scatter', {'x': 'x', 'y': 'y'})
        with self.assertRaisesRegex(ValueError, '不存在'):
            render_figure(spec, [])
        spec['source_ids'] = []
        with self.assertRaisesRegex(ValueError, '先选择'):
            render_figure(spec, [source('x,y\n1,2')])

    def test_log_axis_does_not_silently_drop_nonpositive_points(self):
        spec = recipe('scatter', {'x': 'x', 'y': 'y'})
        spec['style']['xscale'] = 'log'
        with self.assertRaisesRegex(ValueError, '零或负'):
            render_figure(spec, [source('x,y\n0,1\n1,2')])
        spec['style']['xlim'] = [0, 10]
        with self.assertRaisesRegex(ValueError, '范围'):
            validate_spec(spec)

    def test_categorical_tick_explosion_is_rejected(self):
        data = source('condition,score\n' + '\n'.join(f'category{i},{i}' for i in range(90)))
        with self.assertRaisesRegex(ValueError, '80'):
            render_figure(recipe('line', {'x': 'condition', 'y': 'score'}), [data])

    def test_enlarged_composition_title_wraps_to_physical_font_width(self):
        from app.plot_engine import _title_wrap
        text='边缘推理模型压缩与能耗评估 合成验收 错误率'
        wrapped=_title_wrap(text,{'width':6,'font_size':22})
        self.assertGreater(len(wrapped.splitlines()),1)
        self.assertEqual(wrapped.replace('\n','').replace(' ',''),text.replace(' ',''))


if __name__ == '__main__':
    unittest.main()
