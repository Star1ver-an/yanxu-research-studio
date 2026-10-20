"""Figure contracts complement the render-based pagination acceptance tool."""
import io
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from app.word_export import make_docx


def png(width=900, height=500):
    stream=io.BytesIO()
    Image.new('RGB',(width,height),'white').save(stream,format='PNG')
    return stream.getvalue()


class WordFigureTests(unittest.TestCase):
    def figure(self, identity='fig1', **kwargs):
        return dict(id=identity,title='模型性能比较',png=png(),caption='准确率与推理时延的权衡。',**kwargs)

    def test_backward_compatible_report_and_table(self):
        document=Document(io.BytesIO(make_docx('# 研究方案\n\n| 模型 | 准确率 |\n|---|---|\n| 基线 | 0.82 |')))
        self.assertEqual(document.paragraphs[0].style.name,'Title')
        self.assertEqual(len(document.tables),1)
        self.assertEqual(len(document.inline_shapes),0)
        self.assertIsNotNone(document.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')))

    def test_append_png_caption_and_alt_text(self):
        document=Document(io.BytesIO(make_docx('# 研究报告\n测试正文。',[self.figure()])))
        self.assertEqual(len(document.inline_shapes),1)
        text='\n'.join(p.text for p in document.paragraphs)
        self.assertIn('研究图表',text)
        self.assertIn('图 1  模型性能比较。准确率与推理时延的权衡。',text)
        shape=document.inline_shapes[0]
        self.assertIn('模型性能比较',shape._inline.docPr.get('descr'))
        image_paragraph=next(p for p in document.paragraphs if p._p.xpath('.//w:drawing'))
        self.assertTrue(image_paragraph.paragraph_format.keep_with_next)
        self.assertEqual(image_paragraph.paragraph_format.line_spacing,1.0)
        caption=next(p for p in document.paragraphs if p.style.name=='Caption')
        self.assertTrue(caption.paragraph_format.keep_together)
        self.assertFalse(caption.paragraph_format.keep_with_next)
        self.assertFalse(document.styles['Caption'].font.bold)

    def test_place_markers_then_cross_reference_and_missing(self):
        document=Document(io.BytesIO(make_docx('# 方案\n前文 [figure:fig1] 后文\n[figure:fig1]\n[figure:unknown]',[self.figure()])))
        texts=[p.text for p in document.paragraphs]
        self.assertEqual(len(document.inline_shapes),1)
        self.assertLess(texts.index('前文'),texts.index('图 1  模型性能比较。准确率与推理时延的权衡。'))
        self.assertLess(texts.index('图 1  模型性能比较。准确率与推理时延的权衡。'),texts.index('后文'))
        self.assertIn('（见图 1）',texts)
        self.assertIn('〔图表缺失：unknown，请重新关联后导出〕',texts)
        self.assertNotIn('研究图表',texts)

    def test_engine_caption_title_is_not_repeated(self):
        figure=self.figure();figure['caption']='模型性能比较。由本机绘图工具绘制；方法由作者核对。'
        document=Document(io.BytesIO(make_docx('# 研究报告',[figure])))
        caption=next(p.text for p in document.paragraphs if p.style.name=='Caption')
        self.assertEqual(caption.count('模型性能比较'),1)
        self.assertTrue(caption.startswith('图 1  模型性能比较。由本机'))

    def test_provenance_is_retained_in_accessible_image_description(self):
        source_hash='a'*64
        figure=self.figure(provenance='配方版本 2；来源 SHA-256 '+source_hash)
        document=Document(io.BytesIO(make_docx('# 研究报告',[figure])))
        description=document.inline_shapes[0]._inline.docPr.get('descr')
        self.assertIn('来源 SHA-256 '+source_hash,description)
        self.assertIn('配方版本 2',description)
        # Provenance belongs to the picture, keeping the manuscript readable.
        self.assertNotIn(source_hash,'\n'.join(p.text for p in document.paragraphs))

    def test_code_markers_are_literal_and_remaining_append_in_order(self):
        document=Document(io.BytesIO(make_docx('```\n[figure:fig2]\n```',[self.figure(),self.figure('fig2')])))
        text='\n'.join(p.text for p in document.paragraphs)
        self.assertIn('[figure:fig2]',text)
        self.assertEqual(len(document.inline_shapes),2)
        self.assertLess(text.index('图 1'),text.index('图 2'))

    def test_portrait_and_panorama_keep_aspect_and_fit_page(self):
        figures=[dict(id='tall',png=png(500,3500),title='竖图'),dict(id='wide',png=png(3500,500),title='横图')]
        document=Document(io.BytesIO(make_docx('# 研究图表',figures)))
        section=document.sections[0]
        for image,ratio in zip(document.inline_shapes,[1/7,7]):
            self.assertLessEqual(image.width,section.page_width-section.left_margin-section.right_margin)
            self.assertLess(image.height,section.page_height-section.top_margin-section.bottom_margin)
            self.assertAlmostEqual(image.width/image.height,ratio,places=4)

    def test_untrusted_paths_urls_and_invalid_png_are_rejected(self):
        for value in ['C:/secret.png','https://example.com/image.png',b'bad',png()[:-20]]:
            with self.subTest(value_type=type(value).__name__):
                with self.assertRaises(ValueError):
                    make_docx('报告',[dict(id='fig1',png=value)])
        with patch('PIL.Image.open',side_effect=AssertionError('Markdown must not open image files')):
            document=Document(io.BytesIO(make_docx('![图](C:/secret.png) ![远程](https://example.com/image.png)')))
            self.assertEqual(len(document.inline_shapes),0)
            self.assertIn('C:/secret.png',document.paragraphs[0].text)

    def test_duplicate_ids_oversized_images_and_long_captions_reject(self):
        with self.assertRaisesRegex(ValueError,'重复'):
            make_docx('报告',[self.figure(),self.figure()])
        with patch('app.word_export.MAX_PNG_BYTES',3):
            with self.assertRaisesRegex(ValueError,'20 MB'):
                make_docx('报告',[self.figure()])
        with patch('app.word_export.MAX_PNG_PIXELS',2):
            with self.assertRaisesRegex(ValueError,'尺寸'):
                make_docx('报告',[self.figure()])
        with patch('app.word_export.MAX_DOCUMENT_PNG_BYTES',3):
            with self.assertRaisesRegex(ValueError,'分批'):
                make_docx('报告',[self.figure()])
        figure=self.figure();figure['caption']='a'*2001
        with self.assertRaisesRegex(ValueError,'过长'):
            make_docx('报告',[figure])


if __name__=='__main__':
    unittest.main()
