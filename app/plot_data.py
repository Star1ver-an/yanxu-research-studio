"""Bounded, offline readers for plotting inputs; originals are never rewritten."""
import csv
import hashlib
import io
import math
from pathlib import Path
import zipfile

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 200
MAX_PIXELS = 24_000_000
MISSING = {'', 'na', 'n/a', 'null', 'none', 'nan'}
TABLE_SUFFIXES = {'.csv', '.tsv', '.xlsx'}
IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg'}


def _input(raw, name):
    if not isinstance(raw, bytes) or not raw:
        raise ValueError('绘图材料为空或不是有效文件。')
    if len(raw) > MAX_BYTES:
        raise ValueError('单个绘图材料不得超过 20 MB。')
    if not isinstance(name, str) or len(name) > 250:
        raise ValueError('绘图材料名称无效。')
    return Path(name).suffix.lower()


def _cell(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    value = str(value).strip()
    if len(value) > 10000:
        raise ValueError('单个表格单元格过长（最多 10,000 个字符）。')
    if value.lower() in MISSING:
        return None
    # Preserve identifier spelling and integers that cannot survive a JS numeric
    # preview round trip. These remain usable as labels, never silently rounded.
    integer = value.lstrip('+-')
    if integer.isdigit():
        if len(integer) > 1 and integer.startswith('0') or len(integer) > 15:
            return value
        return int(value)
    try:
        number = float(value)
    except ValueError:
        return value
    return number if math.isfinite(number) else None


def _xlsx_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            expanded = sum(x.file_size for x in entries)
            if len(entries) > 2000 or expanded > 100 * 1024 * 1024:
                raise ValueError('Excel 文件解压后过大，请只保留需要的数据表。')
            if any(x.flag_bits & 1 for x in entries):
                raise ValueError('暂不支持加密 Excel，请导出无密码的数值表。')
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError('Excel 文件无效，请使用 .xlsx 文件，或导出 CSV/TSV。') from None
    try:
        if not book.worksheets:
            raise ValueError('Excel 文件没有工作表。')
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ValueError('绘图表格最多支持 50,000 行，请拆分数据。')
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ValueError('绘图表格最多支持 200 列，请删除多余空白列。')
        rows = []
        for row in sheet.iter_rows():
            if len(rows) > MAX_ROWS:
                raise ValueError('绘图表格最多支持 50,000 行。')
            if any(c.data_type == 'f' for c in row):
                raise ValueError('Excel 含有公式；绘图不会执行公式，请先复制并粘贴为值后导入。')
            values = [c.value for c in row]
            if any(v is not None and str(v).strip() for v in values):
                rows.append(values)
        return rows, sheet.title, (['Excel 含多个工作表，本次仅使用第一张：' + sheet.title] if len(book.worksheets) > 1 else [])
    finally:
        book.close()


def load_table(raw, name):
    """Return columns/rows plus bounded metadata, never evaluate cell formulas."""
    suffix = _input(raw, name)
    warnings = []
    sheet = None
    if suffix == '.xlsx':
        records, sheet, warnings = _xlsx_rows(raw)
    elif suffix in {'.csv', '.tsv'}:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = raw.decode('gb18030')
                warnings.append('材料按 GB18030 编码读取；建议使用 UTF-8 保存。')
            except UnicodeDecodeError:
                raise ValueError('表格编码无效，请另存为 UTF-8 CSV/TSV。') from None
        records = []
        try:
            for row in csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ',', strict=True):
                if any(x.strip() for x in row):
                    if len(row) > MAX_COLUMNS:
                        raise ValueError('绘图表格最多支持 200 列。')
                    records.append(row)
                    if len(records) > MAX_ROWS + 1:
                        raise ValueError('绘图表格最多支持 50,000 行。')
        except csv.Error:
            raise ValueError('表格格式无效，请检查引号、分隔符和过长单元格。') from None
    else:
        raise ValueError('绘图数据支持 CSV、TSV 或 XLSX。')
    if len(records) < 2:
        raise ValueError('表格需要一行列名和至少一行数据。')
    header = [str(x or '').strip() for x in records[0]]
    if not all(header) or len(set(header)) != len(header):
        raise ValueError('表格列名不能为空或重复，请先更正。')
    if len(header) > MAX_COLUMNS or any(len(x) > 200 for x in header):
        raise ValueError('表格最多 200 列，列名最多 200 个字符。')
    if any(len(row) != len(header) for row in records[1:]):
        raise ValueError('表格行列数不一致，请检查缺少或多余的分隔符。')
    rows = [dict(zip(header, [_cell(v) for v in row])) for row in records[1:]]
    columns = []
    missing_total = 0
    for field in header:
        values = [r[field] for r in rows if r[field] is not None]
        numeric = sum(isinstance(v, (int, float)) for v in values)
        missing = len(rows) - len(values)
        missing_total += missing
        columns.append({'name': field, 'type': 'numeric' if values and numeric == len(values) else 'mixed' if numeric else 'text',
                        'missing': missing, 'observed': len(values), 'unique': len(set(values))})
    if missing_total:
        warnings.append(f'有 {missing_total} 个空白、缺失标记或非有限数值单元格；绘图会按所需字段排除并报告。')
    return {'columns': columns, 'rows': rows, 'row_count': len(rows), 'warnings': warnings,
            'sha256': hashlib.sha256(raw).hexdigest(), 'sheet': sheet}


def load_image(raw, name):
    """Load a bounded raster copy; never write or alter source bytes."""
    suffix = _input(raw, name)
    if suffix not in IMAGE_SUFFIXES:
        raise ValueError('原图拼版支持 PNG、JPG、JPEG。')
    from PIL import Image, ImageOps
    try:
        with Image.open(io.BytesIO(raw)) as source:
            if source.format not in {'PNG', 'JPEG'}:
                raise ValueError('文件内容与支持的 PNG/JPEG 格式不匹配。')
            if getattr(source, 'n_frames', 1) != 1:
                raise ValueError('暂不支持动画或多帧图片，请选择需要的单帧原图后导入。')
            if source.width * source.height > MAX_PIXELS or max(source.size) > 20000:
                raise ValueError('图片最多支持 2,400 万像素、单边不超过 20,000 像素。')
            orientation = source.getexif().get(274, 1)
            source.load()
            result = ImageOps.exif_transpose(source).convert('RGBA')
            result.info['source_format'] = source.format
            result.info['exif_orientation'] = orientation
            return result
    except ValueError:
        raise
    except Exception:
        raise ValueError('图片损坏或格式无效，请重新导入 PNG/JPEG 原图。') from None


def inspect_source(raw, name):
    suffix = _input(raw, name)
    if suffix in TABLE_SUFFIXES:
        table = load_table(raw, name)
        return {'kind': 'table', 'name': name, 'columns': table['columns'], 'row_count': table['row_count'],
                'preview': table['rows'][:8], 'warnings': table['warnings'], 'sha256': table['sha256'], 'sheet': table['sheet']}
    if suffix in IMAGE_SUFFIXES:
        im = load_image(raw, name)
        return {'kind': 'image', 'name': name, 'columns': [], 'row_count': 0, 'preview': [],
                'width': im.width, 'height': im.height, 'format': im.info['source_format'],
                'sha256': hashlib.sha256(raw).hexdigest(),
                'warnings': (['仅在预览/拼版中按 EXIF 方向显示；原文件保持不变。'] if im.info.get('exif_orientation', 1) != 1 else [])}
    raise ValueError('绘图材料支持 CSV、TSV、XLSX、PNG、JPG、JPEG。')
