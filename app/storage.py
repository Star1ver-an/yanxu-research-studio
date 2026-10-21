import base64
from contextlib import closing, contextmanager
import ctypes
from ctypes import wintypes
from datetime import datetime, timezone
import io
import json
import os
from pathlib import Path
import sqlite3
import threading
import uuid

from .workflows import FLOWS

def now():
    return datetime.now(timezone.utc).isoformat()

def _physical_data_directory(path):
    """Resolve a file, not just its directory, through Windows MSIX redirection.

    A virtual directory can resolve to its logical name while an existing file
    resolves into the package's LocalCache. Probe first-run workspaces as well,
    before the desktop instance lock is taken; never open/recover their store.
    """
    root=Path(path).expanduser().resolve()
    database=root/'workspace.sqlite3'
    if database.is_file():
        return database.resolve().parent
    root.mkdir(parents=True,exist_ok=True)
    probe=root/('.yanxu-location-'+uuid.uuid4().hex+'.tmp')
    try:
        with probe.open('xb') as stream:
            stream.write(b'\0')
        return probe.resolve().parent
    finally:
        probe.unlink(missing_ok=True)


def _workspace_content(root):
    """Read-only discovery: populated, empty, or unreadable (never recover)."""
    database=root/'workspace.sqlite3'
    if not database.is_file():
        return 'empty'
    try:
        with closing(sqlite3.connect(database.resolve().as_uri()+'?mode=ro',uri=True,timeout=.5)) as connection:
            tables={row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if 'projects' in tables and connection.execute('SELECT 1 FROM projects LIMIT 1').fetchone():
                return 'populated'
            if 'settings' in tables:
                for row in connection.execute('SELECT body FROM settings LIMIT 2'):
                    settings=json.loads(row[0])
                    if isinstance(settings,dict) and settings:
                        return 'populated'
        return 'empty'
    except (sqlite3.Error,OSError,ValueError,TypeError):
        return 'unreadable'


def _has_preserved_files(root):
    # Retain an abnormal but potentially valuable workspace with orphaned
    # material/backup files. Runtime logs and lock files do not constitute data.
    try:
        return any(folder.is_dir() and any(folder.iterdir())
                   for folder in (root/'materials',root/'backups'))
    except OSError:
        return True


def data_directory():
    override=os.environ.get('YANXU_DATA_DIR')
    if override:
        return _physical_data_directory(override)
    local=Path(os.environ.get('LOCALAPPDATA',Path.home())).expanduser().resolve()
    root=_physical_data_directory(local/'YanxuResearchStudio')
    if _workspace_content(root)!='empty' or _has_preserved_files(root):
        return root
    candidates=set()
    try:
        package_roots=(local/'Packages').glob('OpenAI.Codex_*/LocalCache/Local/YanxuResearchStudio')
        for candidate in package_roots:
            database=candidate/'workspace.sqlite3'
            if database.is_file():
                physical=database.resolve().parent
                if physical!=root and _workspace_content(physical)=='populated':
                    candidates.add(physical)
    except OSError:
        pass
    # Multiple independent workspaces are never merged or selected arbitrarily.
    return next(iter(candidates)) if len(candidates)==1 else root

class Blob(ctypes.Structure):
    _fields_=[('cbData',wintypes.DWORD),('pbData',ctypes.POINTER(ctypes.c_byte))]

def protect(text, decrypt=False):
    if os.name!='nt':
        raise ValueError('此版本只在 Windows 上保存加密密钥。')
    raw=base64.b64decode(text) if decrypt else text.encode('utf-8')
    buf=ctypes.create_string_buffer(raw)
    source=Blob(len(raw),ctypes.cast(buf,ctypes.POINTER(ctypes.c_byte)))
    dest=Blob()
    fn=ctypes.windll.crypt32.CryptUnprotectData if decrypt else ctypes.windll.crypt32.CryptProtectData
    if not fn(ctypes.byref(source),None,None,None,None,1,ctypes.byref(dest)):
        raise ValueError('Windows 无法处理加密密钥。请重新输入。')
    try:
        out=ctypes.string_at(dest.pbData,dest.cbData)
        return out.decode('utf-8') if decrypt else base64.b64encode(out).decode('ascii')
    finally:
        ctypes.windll.kernel32.LocalFree(dest.pbData)

DEFAULT_SETTINGS=dict(mode='codex',codex_path='',codex_model='',provider='compatible',
                      base_url='https://api.deepseek.com/v1',api_model='deepseek-chat',max_tokens=6000)

class Store:
    def __init__(self,path=None):
        self.root=_physical_data_directory(path) if path else data_directory()
        self.root.mkdir(parents=True,exist_ok=True)
        self.lock=threading.RLock()
        self.db=self.root/'workspace.sqlite3'
        with self.connect() as c:
            c.execute('CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY, body TEXT NOT NULL)')
            c.execute('CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY, body TEXT NOT NULL)')
            c.execute('CREATE TABLE IF NOT EXISTS preferences(key TEXT PRIMARY KEY, body TEXT NOT NULL)')
        self.recover()

    @contextmanager
    def connect(self):
        connection=sqlite3.connect(self.db,timeout=15)
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def get_settings(self,secret=False):
        with self.connect() as c:
            row=c.execute('SELECT body FROM settings WHERE id=1').fetchone()
        s={**DEFAULT_SETTINGS,**(json.loads(row[0]) if row else {})}
        encrypted=s.pop('encrypted_key','')
        if secret:
            s['api_key']=protect(encrypted,True) if encrypted else ''
        else:
            s['has_key']=bool(encrypted)
        return s

    def save_settings(self,data):
        with self.lock:
            with self.connect() as c:
                row=c.execute('SELECT body FROM settings WHERE id=1').fetchone()
                s={**DEFAULT_SETTINGS,**(json.loads(row[0]) if row else {})}
                for k in DEFAULT_SETTINGS:
                    if k in data:
                        s[k]=data[k]
                if s['mode'] not in ('codex','api') or s['provider'] not in ('compatible','anthropic'):
                    raise ValueError('不支持的连接类型。')
                s['max_tokens']=max(256,min(32000,int(s['max_tokens'])))
                if data.get('clear_key'):
                    s.pop('encrypted_key',None)
                elif data.get('api_key','').strip():
                    s['encrypted_key']=protect(data['api_key'].strip())
                c.execute('INSERT OR REPLACE INTO settings VALUES(1,?)',(json.dumps(s),))
        return self.get_settings()

    def list(self):
        with self.connect() as c:
            values=[json.loads(r[0]) for r in c.execute('SELECT body FROM projects')]
        return sorted(values,key=lambda p:p['updated'],reverse=True)

    def language(self):
        with self.connect() as connection:
            row=connection.execute("SELECT body FROM preferences WHERE key='language_v1'").fetchone()
        return {'ui_language':'zh-CN',**(json.loads(row[0]) if row else {})}

    def save_language(self,data):
        from .i18n import language
        value={'ui_language':language(data.get('ui_language'))}
        with self.lock,self.connect() as connection:
            connection.execute('INSERT OR REPLACE INTO preferences VALUES(?,?)',('language_v1',json.dumps(value)))
        return value

    def onboarding(self):
        with self.connect() as c:
            row=c.execute('SELECT body FROM preferences WHERE key=?',('onboarding_v1',)).fetchone()
        return json.loads(row[0]) if row else {'status':'not_started','step':0}

    def save_onboarding(self,data):
        status=data.get('status');step=data.get('step')
        if status not in ('in_progress','completed','dismissed') or type(step)!=int or not 0<=step<6:
            raise ValueError('入门引导进度无效。')
        state={'status':status,'step':step}
        with self.lock,self.connect() as c:
            c.execute('INSERT OR REPLACE INTO preferences VALUES(?,?)',('onboarding_v1',json.dumps(state)))
        return state

    def get(self,pid):
        with self.connect() as c:
            row=c.execute('SELECT body FROM projects WHERE id=?',(pid,)).fetchone()
        if not row:
            raise ValueError('未找到这个项目。')
        return json.loads(row[0])

    def put(self,p):
        p['updated']=now()
        with self.lock,self.connect() as c:
            c.execute('INSERT OR REPLACE INTO projects VALUES(?,?)',(p['id'],json.dumps(p,ensure_ascii=False)))
        return p

    def create(self,data):
        from .i18n import language
        output=data.get('output_language','auto')
        if output!='auto':language(output)
        flow=data.get('flow','full')
        if flow not in FLOWS:
            raise ValueError('请选择有效的研究流程。')
        route=data.get('route',FLOWS[flow]['route'])
        if route not in ('ars','scientific') or route!=FLOWS[flow]['route']:
            raise ValueError('研究路径与工作流不匹配，请重新选择。')
        selected_skill=str(data.get('selected_skill',''))
        if selected_skill:
            if route!='scientific':raise ValueError('专业技能只能用于科研工具与分析路径。')
            from .scientific import skill_text
            skill_text(selected_skill,0)
        if flow=='scientific_task' and not selected_skill:
            raise ValueError('请先为专业技能专题选择一项技能。')
        from .routing import recommend
        recommendation=recommend(data) if data.get('route_choice')=='recommended' else None
        if recommendation and (recommendation['route']!=route or recommendation['flow']!=flow):
            raise ValueError('推荐与当前目标或选择不一致，请重新获取推荐或手动选择路径。')
        title=str(data.get('title','')).strip()[:140]
        if not title:
            raise ValueError('请填写项目名称。')
        declaration=data.get('experiments','none')
        if declaration not in ('none','declared','undecided'):
            raise ValueError('请选择实验来源声明。')
        p=dict(id=uuid.uuid4().hex,title=title,topic=str(data.get('topic',''))[:20000],
               discipline=str(data.get('discipline','跨学科'))[:200],flow=flow,route=route,
               selected_skill=selected_skill,goal=str(data.get('goal','unsure'))[:30],
               data_kind=str(data.get('data_kind','unknown'))[:30],
               route_choice='recommended' if data.get('route_choice')=='recommended' else 'manual',
               routing_recommendation=recommendation,
               citation=str(data.get('citation','APA 7'))[:100],experiments=declaration,
               requirements=str(data.get('requirements',''))[:20000],created=now(),updated=now(),
               cursor=0,assets=[],literature=[],runs=[],outputs={},confirmations=[],demo=False,output_language=output)
        return self.put(p)

    def recover(self):
        for p in self.list():
            changed=False
            for r in p.get('runs',[]):
                if r['status']=='running':
                    r.update(status='interrupted',error='上次运行随软件关闭而中断，可重新运行。')
                    changed=True
            if changed:
                self.put(p)

def extract_file(name,raw):
    suffix=Path(name).suffix.lower()
    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('该 PDF 已加密，请先解密后导入。')
        if len(reader.pages)>600:
            raise ValueError('PDF 超过 600 页，请拆分后导入。')
        text='\n\n'.join(f'[第 {i+1} 页]\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages))
        if len(text.strip())<35:
            raise ValueError('未提取到可用文本，这可能是扫描版 PDF。请先 OCR 或导入文字版。')
    elif suffix=='.docx':
        from docx import Document
        doc=Document(io.BytesIO(raw))
        parts=[p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [' | '.join(cell.text for cell in row.cells) for row in table.rows]
        text='\n'.join(parts)
    elif suffix in ('.txt','.md','.csv','.ris','.nbib','.bib','.json','.tsv'):
        try:
            text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text=raw.decode('gb18030')
    elif suffix in ('.h5ad','.zip','.gmt','.fcs','.vcf'):
        from .analysis_special import inspect_special
        info=inspect_special(raw,name)
        text='本机专用分析材料（仅结构摘要，原始记录不自动发送给模型）：'+json.dumps({k:v for k,v in info.items() if k not in ('preview','rows','data')},ensure_ascii=False)
    elif suffix in ('.xlsx','.png','.jpg','.jpeg'):
        from .plot_data import inspect_source
        info=inspect_source(raw,name)
        # Binary plotting materials remain local. No data rows or pixels in stage context.
        text='本机绘图材料（未读取为写作正文）：'+json.dumps({k:v for k,v in info.items() if k!='preview'},ensure_ascii=False)
    else:
        raise ValueError('支持 PDF、DOCX、TXT、Markdown、CSV、TSV、XLSX、PNG、JPG、RIS、NBIB、BibTeX、JSON，以及分析材料 H5AD、10x MTX ZIP、GMT、FCS、VCF。')
    if not text.strip():
        raise ValueError('文件没有可用文本。')
    if len(text)>1200000:
        raise ValueError('提取文本过长，请拆分文件后导入。')
    return text
